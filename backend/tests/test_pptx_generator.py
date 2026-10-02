# -*- coding: utf-8 -*-
"""PPT 导出：文本清洗 + `_add_bullets` 只能有一份定义。

这个文件补的是一块**此前完全没有测试覆盖**的地方（`backend/tests` 里搜不到任何 pptx
用例），而它出过两次事：

1. **同名函数覆盖。** `_add_bullets` 曾经在同一个文件里写了三遍，后定义的静默盖掉前面的
   （Python 不报错、不警告），于是"哪一份在跑"只由文件里的先后顺序决定 —— 而最后那份恰好
   是最弱的：不清理文本、不做字号自适应、**丢掉关键术语加粗**。已经合并成一份，这里把它钉住。
2. **markdown 标记漏进成品。** 生成侧没有 `**` 这个约定（见 `prompts/resource/ppt.yaml`
   的格式说明），是模型自己加的强调；渲染侧原先没有任何地方处理它，于是导出的 PPT 里
   课文原样带着 `**组件定义**` 这种星号。实测一份真实资源导出 19 页、正文里就带着星号。

第 2 条的修法有个陷阱：`**` 也是 Python 的幂运算符。剥标记的正则要求 `**` 两侧紧贴非空白，
所以 `a ** b ** c` 不会被误伤 —— 这条也有用例守着，因为**改坏代码比留下星号严重得多**。
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from backend.src.utils import pptx_generator as pptx  # noqa: E402
from backend.src.utils.pptx_generator import _clean_ppt_text, _strip_emphasis  # noqa: E402


# ── 只能有一份定义 ──

def test_add_bullets_is_defined_exactly_once():
    """同名函数写两遍不会报错，只会**静默地**让其中一份永不生效。

    这条测试不看行为、只看源码里 `def _add_bullets` 出现几次 —— 行为测试挡不住这个坑：
    三份定义里注释掉两份、留下的那份行为正确，所有行为测试照样全绿，而"读代码的人以为
    自己改的那份在跑"这件事没有任何症状。
    """
    source = Path(pptx.__file__).read_text(encoding="utf-8")
    assert len(re.findall(r"^def _add_bullets\b", source, flags=re.MULTILINE)) == 1


# ── 加粗标记 ──

def test_emphasis_markers_are_stripped_from_the_text():
    assert _strip_emphasis("**组件定义**：智能体 = 大模型 + 工具") == "组件定义：智能体 = 大模型 + 工具"
    assert _strip_emphasis("普通文字没有标记") == "普通文字没有标记"


def test_emphasis_inside_a_sentence_is_stripped_too():
    assert _strip_emphasis("缺**工具**只能聊天") == "缺工具只能聊天"


def test_python_exponentiation_is_left_alone():
    """`a ** b ** c` 是幂运算，不是强调。宽松的正则会把中间咬掉，**把代码改坏**。"""
    assert _strip_emphasis("x = a ** b ** c") == "x = a ** b ** c"
    assert _strip_emphasis("2 ** 10 等于 1024") == "2 ** 10 等于 1024"


def test_clean_text_strips_markers_as_well():
    """`_clean_ppt_text` 是所有文本进 PPT 的必经之路（标题、副标题、要点、步骤、
    公式、备注都走它），所以标记必须在这一层就没了。"""
    assert "**" not in _clean_ppt_text("**核心规则**：默认值在定义时求值一次")
    assert _clean_ppt_text("**核心规则**：默认值在定义时求值一次").startswith("核心规则：")


def test_clean_text_still_cleans_html_and_comments():
    """改动不能挤掉原有的清洗职责。"""
    cleaned = _clean_ppt_text("<!-- layout: content_cards -->正文<span>x</span>")
    assert "layout" not in cleaned
    assert "<span>" not in cleaned
    assert "正文" in cleaned


def test_clean_text_honours_the_length_limit():
    """`limit` 管的是**留下的字数**（`limit - 1` 个字符），省略号加在它外面 ——
    所以结果长度是 `limit - 1 + 3`。这是原有行为，顺手钉住，免得日后误以为它等于 limit。"""
    cleaned = _clean_ppt_text("字" * 50, 10)
    assert cleaned == "字" * 9 + "..."
    assert len(cleaned) == 10 - 1 + 3


# ── 真实形状的端到端：一份真实资源里的那行 ──

def test_a_real_bullet_line_renders_without_stars_and_with_a_bold_key():
    """抄自库里一份真实 PPT 正文（`generated_resources` id=332）里的一行。

    它同时验证两件事：星号进不了 PPT；「中文冒号前加粗成深蓝」这一步还在
    （合并三份定义时特意保留的能力，丢过一次）。
    """
    from pptx import Presentation

    line = "- **组件定义**：智能体 = 大模型（推理与决策）+ 工具（行动接口）"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pptx._add_bullets(slide, [line])

    paragraph = slide.shapes[-1].text_frame.paragraphs[0]
    assert "**" not in paragraph.text, "星号漏进了成品"
    bold_runs = [run.text for run in paragraph.runs if run.font.bold]
    assert bold_runs, "关键术语没有被加粗（合并定义时丢过一次的能力）"
